"""W-20 agreement between two raters, and W-21 extraction against the gold standard.

    python scripts/w20_agreement.py rater1.xlsx rater2.xlsx \
        [--pilot data/curriculum/w21_pilot_v0.jsonl] [--out results/w20-agreement.md]

Reads the "評分" sheet written by scripts/make_w20_rating_sheet.py.

PRE-REGISTERED (rubric.md §4.3, committed before any rating exists):
  - Rater agreement per dimension: quadratic-weighted Cohen's kappa on the five
    ordered levels {0, 0.25, 0.5, 0.75, 1}; pairs where either rater wrote "缺值"
    (missing) are excluded and counted. kappa >= 0.70 -> dimension accepted; else the
    anchors of that dimension are revised and re-rated (rubric step 3).
  - Cases with |difference| >= 0.5 are listed for third-party adjudication (step 4).
  - Gold standard = the mean of the two raters where |difference| <= 0.25; adjudicated
    cases are added later; other cases are left out of the gold standard.
  - W-21 against gold (V-C0-2): agreement = |extraction - gold| <= 0.25; a missing
    extraction where gold exists counts as disagreement. Pass >= 90% per dimension.
  - Missing-value agreement (both "缺值") is reported separately; it is not kappa.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np

LEVELS = (0.0, 0.25, 0.5, 0.75, 1.0)
DIMS = ("quant_intensity", "programming_intensity", "research_orientation", "entrepreneurship")
MISSING = "缺值"
KAPPA_PASS = 0.70
V_C0_2_PASS = 0.90


def parse_score(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", MISSING):
        return None
    x = float(s)
    if x not in LEVELS:
        raise ValueError(f"score {s!r} is not one of {LEVELS} or {MISSING}")
    return x


def read_ratings(path: Path) -> dict[tuple[str, str], float | None]:
    """{(program code, dimension): score or None}; blank rows are rejected."""
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)["評分"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    i_id, i_dim, i_score = head.index("編號"), head.index("維度代碼"), head.index("分數")
    out: dict[tuple[str, str], float | None] = {}
    blank = []
    for r in rows[1:]:
        if r[i_id] is None:
            continue
        key = (str(r[i_id]), str(r[i_dim]))
        if r[i_score] is None or str(r[i_score]).strip() == "":
            blank.append(key)
            continue
        out[key] = parse_score(r[i_score])
    if blank:
        raise ValueError(f"{path}: {len(blank)} unrated rows, e.g. {blank[:3]} (use {MISSING})")
    return out


def weighted_kappa(a: list[float], b: list[float]) -> float:
    """Quadratic-weighted Cohen's kappa over LEVELS (Cohen 1968)."""
    if len(a) != len(b) or not a:
        raise ValueError("need two equal-length, non-empty rating lists")
    k = len(LEVELS)
    ia = np.array([LEVELS.index(x) for x in a])
    ib = np.array([LEVELS.index(x) for x in b])
    obs = np.zeros((k, k))
    np.add.at(obs, (ia, ib), 1)
    obs /= obs.sum()
    exp = np.outer(obs.sum(axis=1), obs.sum(axis=0))
    w = np.array([[(i - j) ** 2 for j in range(k)] for i in range(k)]) / (k - 1) ** 2
    den = (w * exp).sum()
    if den == 0:  # both raters constant (and equal): kappa undefined
        return float("nan")
    return float(1 - (w * obs).sum() / den)


def _mean(xs: list[bool]) -> float:
    return float(np.mean(xs)) if xs else float("nan")


def agreement(r1: dict, r2: dict) -> dict:
    keys = sorted(set(r1) | set(r2))
    missing_keys = sorted(set(r1) ^ set(r2))
    if missing_keys:
        raise ValueError(f"the two sheets rate different rows, e.g. {missing_keys[:3]}")
    res = {}
    for d in DIMS:
        pairs = [(k, r1[k], r2[k]) for k in keys if k[1] == d]
        both = [(k, x, y) for k, x, y in pairs if x is not None and y is not None]
        a, b = [x for _, x, _ in both], [y for _, _, y in both]
        kappa = weighted_kappa(a, b) if both else float("nan")
        res[d] = {
            "n": len(pairs),
            "n_scored_both": len(both),
            "n_missing_both": sum(x is None and y is None for _, x, y in pairs),
            "n_missing_one": sum((x is None) != (y is None) for _, x, y in pairs),
            "kappa": kappa,
            "exact": _mean([x == y for x, y in zip(a, b, strict=True)]),
            "within_025": _mean([abs(x - y) <= 0.25 for x, y in zip(a, b, strict=True)]),
            "accepted": bool(kappa >= KAPPA_PASS) if both else False,
            "adjudicate": [k[0] for k, x, y in both if abs(x - y) >= 0.5],
            "gold": {k[0]: (x + y) / 2 for k, x, y in both if abs(x - y) <= 0.25},
        }  # fmt: skip
    return res


def extraction_vs_gold(res: dict, pilot: list[dict]) -> dict:
    """V-C0-2 on programmes whose order in the pilot file defines W01..W30."""
    code = {f"W{i:02d}": rec for i, rec in enumerate(pilot, 1)}
    out = {}
    for d in DIMS:
        gold = res[d]["gold"]
        hits = []
        for pid, g in gold.items():
            v = code[pid]["scores"][d]["value"] if pid in code else None
            hits.append(v is not None and abs(float(v) - g) <= 0.25)
        rate = float(np.mean(hits)) if hits else float("nan")
        out[d] = {
            "n_gold": len(gold),
            "agreement": rate,
            "pass": bool(hits) and rate >= V_C0_2_PASS,
        }
    return out


def report(res: dict, ext: dict | None) -> str:
    L = ["# W-20 評分者一致性（預先登記：scripts/w20_agreement.py）", "",
         "| 維度 | 兩人皆有分數 | 兩人皆缺值 | 一人缺值 | 加權 κ | 完全一致 "
         "| 差距 ≤ 0.25 | 判定（κ ≥ 0.70） | 需裁決 |",
         "|---|---|---|---|---|---|---|---|---|"]  # fmt: skip
    for d, r in res.items():
        verdict = "✅ 接受" if r["accepted"] else "❌ 修訂錨點後重評"
        L.append(f"| {d} | {r['n_scored_both']} | {r['n_missing_both']} | {r['n_missing_one']} | "
                 f"{r['kappa']:.2f} | {r['exact']:.0%} | {r['within_025']:.0%} | {verdict} | "
                 f"{'、'.join(r['adjudicate']) or '—'} |")  # fmt: skip
    if ext is not None:
        L += ["", "## W-21 擷取 vs 黃金標準（V-C0-2，門檻 90%）", "",
              "| 維度 | 黃金標準筆數 | 一致率（差距 ≤ 0.25） | 判定 |",
              "|---|---|---|---|"]  # fmt: skip
        for d, e in ext.items():
            L.append(
                f"| {d} | {e['n_gold']} | {e['agreement']:.0%} | {'✅' if e['pass'] else '❌'} |"
            )
    return "\n".join(L) + "\n"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("rater1", type=Path)
    ap.add_argument("rater2", type=Path)
    ap.add_argument("--pilot", type=Path, default=None, help="W-21 JSONL to score against gold")
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()
    res = agreement(read_ratings(args.rater1), read_ratings(args.rater2))
    ext = None
    if args.pilot:
        pilot = [json.loads(x) for x in args.pilot.read_text(encoding="utf-8").splitlines() if x]
        ext = extraction_vs_gold(res, pilot)
    md = report(res, ext)
    if args.out:
        args.out.write_text(md, encoding="utf-8")
    print(md)


if __name__ == "__main__":
    main()
