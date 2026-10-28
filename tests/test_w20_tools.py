"""W-20 rating workbook and pre-registered agreement script (L2)."""

import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PILOT = ROOT / "data" / "curriculum" / "w21_pilot_v0.jsonl"


def _load(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


make = _load("make_w20_rating_sheet")
agr = _load("w20_agreement")
RECS = [json.loads(x) for x in PILOT.read_text(encoding="utf-8").splitlines() if x]


def _sheet(tmp_path, rater, fill):
    path = tmp_path / f"{rater}.xlsx"
    make.build(RECS, rater, path)
    wb = load_workbook(path)
    ws = wb["評分"]
    head = [c.value for c in ws[1]]
    i_id, i_dim, i_score = head.index("編號"), head.index("維度代碼"), head.index("分數")
    for row in ws.iter_rows(min_row=2):
        if row[i_id].value:
            row[i_score].value = fill(row[i_id].value, row[i_dim].value)
    wb.save(path)
    return path


def test_kappa_known_values():
    assert agr.weighted_kappa([0, 0.5, 1, 1], [0, 0.5, 1, 1]) == pytest.approx(1.0)
    # hand computation: two raters, two levels used (0 and 1), one disagreement in four
    a, b = [0.0, 0.0, 1.0, 1.0], [0.0, 1.0, 1.0, 1.0]
    obs_dis = 1 / 4  # weight 1 for the 0-vs-1 cell (max distance)
    exp_dis = (2 / 4) * (1 / 4) + (2 / 4) * (3 / 4)  # P(a=0)P(b=1) + P(a=1)P(b=0)
    assert agr.weighted_kappa(a, b) == pytest.approx(1 - obs_dis / exp_dis)
    assert agr.weighted_kappa([0, 1, 0, 1], [1, 0, 1, 0]) < 0


def test_kappa_matches_sklearn():
    sk = pytest.importorskip("sklearn.metrics")
    a = [0, 0.25, 0.5, 0.5, 0.75, 1, 1, 0, 0.25, 0.5]
    b = [0, 0.5, 0.5, 0.25, 0.75, 0.75, 1, 0.25, 0.25, 1]
    ref = sk.cohen_kappa_score([int(x * 4) for x in a], [int(x * 4) for x in b],
                               weights="quadratic", labels=list(range(5)))  # fmt: skip
    assert agr.weighted_kappa(a, b) == pytest.approx(ref)


def test_sheet_is_blind_to_w21_scores(tmp_path):
    path = tmp_path / "R2.xlsx"
    make.build(RECS, "R2", path)
    text = []
    for ws in load_workbook(path).worksheets:
        for row in ws.iter_rows(values_only=True):
            text += [str(v) for v in row if v is not None]
    blob = "\n".join(text)
    # a quote may coincide with the degree name shown on the sheet; that is not a score
    quotes = [s["quote"] for r in RECS for s in r["scores"].values()
              if s.get("quote") and s["quote"] not in r["degree_name"]]  # fmt: skip
    assert quotes and not any(q in blob for q in quotes if len(q) > 12)
    codes = [t for t in text if t.startswith("W") and t[1:].isdigit()]
    assert len(codes) == len(RECS) * 4 + len(RECS)  # rating rows + entry-requirement rows


def test_agreement_end_to_end(tmp_path):
    def r1(pid, dim):
        return "缺值" if dim == "entrepreneurship" and pid == "W01" else 0.5

    def r2(pid, dim):
        if pid == "W02" and dim == "quant_intensity":
            return 1  # differs by 0.5 -> adjudication
        return "缺值" if dim == "entrepreneurship" and pid == "W01" else 0.5

    res = agr.agreement(agr.read_ratings(_sheet(tmp_path, "R1", r1)),
                        agr.read_ratings(_sheet(tmp_path, "R2", r2)))  # fmt: skip
    q = res["quant_intensity"]
    assert q["n"] == len(RECS) and q["adjudicate"] == ["W02"] and "W02" not in q["gold"]
    assert res["entrepreneurship"]["n_missing_both"] == 1
    ext = agr.extraction_vs_gold(res, RECS)
    assert set(ext) == set(agr.DIMS) and 0 <= ext["programming_intensity"]["agreement"] <= 1
    md = agr.report(res, ext)
    assert "加權 κ" in md and "V-C0-2" in md


def test_unrated_rows_and_bad_scores_are_rejected(tmp_path):
    with pytest.raises(ValueError, match="unrated"):
        agr.read_ratings(_sheet(tmp_path, "R1", lambda pid, dim: None if pid == "W03" else 0))
    with pytest.raises(ValueError):
        agr.parse_score("0.3")


def test_manual_matches_rubric():
    """The rater manual copies rubric §4.1-4.2; the rubric is the source of truth."""
    rubric = (ROOT / "docs" / "rubric.md").read_text(encoding="utf-8")
    for row in make.ANCHORS:
        for cell in row[1:]:
            assert (
                cell.replace("非研究型", "**非研究型**").replace("可選", "**可選**") in rubric
                or cell in rubric
            ), cell
